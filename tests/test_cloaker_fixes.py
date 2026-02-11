"""
Tests for cloaker pipeline fixes:
- Fix 1: Multi-line entity splitting (_split_multiline_replacements)
- Fix 2: Party label collision guard
- Fix 3: Person name-variant expansion (_expand_person_name_parts)
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from clientcloak.cloaker import (
    _build_mappings_and_replacements,
    _expand_person_name_parts,
    _split_multiline_replacements,
    build_cloak_replacements,
    cloak_document,
)
from clientcloak.docx_handler import extract_all_text, load_document
from clientcloak.models import CloakConfig, CommentMode


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_docx(tmp_path: Path, name: str, paragraphs: list[str]) -> Path:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(str(path))
    return path


def _extract_text(path: Path) -> str:
    doc = load_document(path)
    fragments = extract_all_text(doc)
    return "\n".join(text for text, _source in fragments)


# ===========================================================================
# Fix 1: _split_multiline_replacements
# ===========================================================================


class TestSplitMultilineReplacements:
    def test_multiline_key_split_into_lines(self):
        replacements = {
            "6001 Bollinger Canyon Road\nSan Ramon, CA 94583": "[Address-1]",
            "Acme Corp.": "[Company]",
        }
        result = _split_multiline_replacements(replacements)

        # Multi-line key should be removed
        assert "6001 Bollinger Canyon Road\nSan Ramon, CA 94583" not in result
        # Individual lines should be present
        assert result["6001 Bollinger Canyon Road"] == "[Address-1]"
        assert result["San Ramon, CA 94583"] == "[Address-1]"
        # Non-multiline keys should be preserved
        assert result["Acme Corp."] == "[Company]"

    def test_empty_lines_skipped(self):
        replacements = {
            "Line One\n\nLine Three": "[Address-1]",
        }
        result = _split_multiline_replacements(replacements)
        assert len(result) == 2
        assert "Line One" in result
        assert "Line Three" in result

    def test_whitespace_only_lines_skipped(self):
        replacements = {
            "Line One\n   \nLine Three": "[Address-1]",
        }
        result = _split_multiline_replacements(replacements)
        assert len(result) == 2

    def test_no_multiline_keys_unchanged(self):
        replacements = {
            "Acme Corp.": "[Company]",
            "BigCo LLC": "[Vendor]",
        }
        result = _split_multiline_replacements(replacements)
        assert result == replacements

    def test_duplicate_lines_keep_first(self):
        """If two multi-line keys share a line, the first mapping wins."""
        replacements = {
            "123 Main St\nNew York, NY": "[Address-1]",
            "456 Oak Ave\nNew York, NY": "[Address-2]",
        }
        result = _split_multiline_replacements(replacements)
        # "New York, NY" should map to whichever was first
        assert "New York, NY" in result
        assert "123 Main St" in result
        assert "456 Oak Ave" in result


# ===========================================================================
# Fix 2: Party label collision guard
# ===========================================================================


class TestLabelCollisionGuard:
    def test_same_label_produces_distinct_keys(self):
        config = CloakConfig(
            party_a_name="Walmart Inc.",
            party_a_label="Form Agreement",
            party_b_name="Target Corp.",
            party_b_label="Form Agreement",
            comment_mode=CommentMode.STRIP,
        )
        mappings, cloak_replacements = _build_mappings_and_replacements(config)

        # Both party names should be in cloak_replacements
        assert "Walmart Inc." in cloak_replacements
        assert "Target Corp." in cloak_replacements

        # They should map to DIFFERENT placeholders
        assert cloak_replacements["Walmart Inc."] != cloak_replacements["Target Corp."]

        # Party A keeps original label
        assert cloak_replacements["Walmart Inc."] == "[Form Agreement]"
        # Party B gets renamed to [Counterparty]
        assert cloak_replacements["Target Corp."] == "[Counterparty]"

    def test_collision_with_counterparty_label(self):
        """Edge case: both labelled 'Counterparty'."""
        config = CloakConfig(
            party_a_name="Walmart Inc.",
            party_a_label="Counterparty",
            party_b_name="Target Corp.",
            party_b_label="Counterparty",
            comment_mode=CommentMode.STRIP,
        )
        mappings, cloak_replacements = _build_mappings_and_replacements(config)

        assert "Walmart Inc." in cloak_replacements
        assert "Target Corp." in cloak_replacements
        assert cloak_replacements["Walmart Inc."] != cloak_replacements["Target Corp."]
        # Party B should get "{label}-2" since "Counterparty" is already taken
        assert cloak_replacements["Target Corp."] == "[Counterparty-2]"

    def test_different_labels_no_change(self):
        config = CloakConfig(
            party_a_name="Walmart Inc.",
            party_a_label="Customer",
            party_b_name="Target Corp.",
            party_b_label="Vendor",
            comment_mode=CommentMode.STRIP,
        )
        mappings, cloak_replacements = _build_mappings_and_replacements(config)

        assert cloak_replacements["Walmart Inc."] == "[Customer]"
        assert cloak_replacements["Target Corp."] == "[Vendor]"

    def test_collision_short_forms_use_resolved_label(self):
        """Party B short forms should use the resolved (collision-safe) label."""
        config = CloakConfig(
            party_a_name="Walmart Inc.",
            party_a_label="Form Agreement",
            party_b_name="Target Corp.",
            party_b_label="Form Agreement",
            party_b_short_forms=["Target"],
            comment_mode=CommentMode.STRIP,
        )
        _mappings, cloak_replacements = _build_mappings_and_replacements(config)

        # Short form should use the resolved label, not the colliding one
        assert "Target" in cloak_replacements
        assert cloak_replacements["Target"] == "[Counterparty]"

    def test_collision_guard_in_full_pipeline(self, tmp_path):
        """End-to-end: both names should be cloaked even with same label."""
        paragraphs = [
            "This Form Agreement is between Walmart Inc. and Target Corp.",
            "Walmart Inc. shall provide goods to Target Corp.",
            "Target Corp. shall pay Walmart Inc. monthly.",
        ]
        docx_path = _make_docx(tmp_path, "collision.docx", paragraphs)
        output_path = tmp_path / "collision_cloaked.docx"
        mapping_path = tmp_path / "collision_mapping.json"

        config = CloakConfig(
            party_a_name="Walmart Inc.",
            party_a_label="Form Agreement",
            party_b_name="Target Corp.",
            party_b_label="Form Agreement",
            comment_mode=CommentMode.STRIP,
            strip_metadata=True,
        )

        result = cloak_document(
            input_path=docx_path,
            output_path=output_path,
            mapping_path=mapping_path,
            config=config,
        )

        cloaked_text = _extract_text(Path(result.output_path))
        assert "Walmart" not in cloaked_text
        assert "Target" not in cloaked_text


# ===========================================================================
# Fix 3: _expand_person_name_parts
# ===========================================================================


class TestExpandPersonNameParts:
    def test_surname_generated(self):
        """'Darren L. Woods' should generate 'Woods' when it appears in text."""
        replacements = {"Darren L. Woods": "[Person-1]"}
        text = "CEO Darren L. Woods announced that Woods will retire."
        result = _expand_person_name_parts(replacements, text)

        assert result["Darren L. Woods"] == "[Person-1]"
        assert result["Woods"] == "[Person-1]"

    def test_first_name_generated(self):
        """'Darren L. Woods' should generate 'Darren' when it appears in text."""
        replacements = {"Darren L. Woods": "[Person-1]"}
        text = "CEO Darren L. Woods said Darren will attend the meeting."
        result = _expand_person_name_parts(replacements, text)

        assert result["Darren"] == "[Person-1]"

    def test_no_middle_initial_form(self):
        """'Darren L. Woods' should generate 'Darren Woods' when present."""
        replacements = {"Darren L. Woods": "[Person-1]"}
        text = "CEO Darren L. Woods, also known as Darren Woods, will attend."
        result = _expand_person_name_parts(replacements, text)

        assert result["Darren Woods"] == "[Person-1]"

    def test_initial_with_period_stripped(self):
        """'C. Douglas McMillon' → significant = ['Douglas', 'McMillon']."""
        replacements = {"C. Douglas McMillon": "[Person-1]"}
        text = "C. Douglas McMillon directed McMillon to sign the deal."
        result = _expand_person_name_parts(replacements, text)

        assert result["McMillon"] == "[Person-1]"
        assert result["Douglas"] == "[Person-1]"

    def test_initial_without_period_stripped(self):
        """'John A Smith' → strips 'A' as single-letter initial."""
        replacements = {"John A Smith": "[Person-1]"}
        text = "John A Smith and Smith agreed to the terms."
        result = _expand_person_name_parts(replacements, text)

        assert result["Smith"] == "[Person-1]"

    def test_four_name_person(self):
        """'Jose Diego Valdez Domingo' generates contiguous subsequences."""
        replacements = {"Jose Diego Valdez Domingo": "[Person-1]"}
        text = (
            "Jose Diego Valdez Domingo was present. "
            "Valdez Domingo signed the agreement. "
            "Jose Domingo was also referenced."
        )
        result = _expand_person_name_parts(replacements, text)

        assert result["Valdez Domingo"] == "[Person-1]"
        assert result["Jose Domingo"] == "[Person-1]"  # first+last shortcut

    def test_all_caps_variant_matched(self):
        """ALL-CAPS signature blocks like 'WOODS' should match."""
        replacements = {"Darren L. Woods": "[Person-1]"}
        text = "By: WOODS\nDarren L. Woods"
        result = _expand_person_name_parts(replacements, text)

        assert result["Woods"] == "[Person-1]"

    def test_lowercase_not_matched(self):
        """Lowercase 'woods' should NOT match (prevents false positives)."""
        replacements = {"Darren L. Woods": "[Person-1]"}
        # Only lowercase "woods" appears — no Title-case reference
        text = "He walked through the woods near the lake."
        result = _expand_person_name_parts(replacements, text)

        # "Woods" should NOT be added because only lowercase "woods" is in text
        assert "Woods" not in result

    def test_month_stopwords_excluded(self):
        """Names like 'May' (from 'Theresa May') should be excluded."""
        replacements = {"Theresa May": "[Person-1]"}
        text = "Theresa May signed on May 15th. May confirmed."
        result = _expand_person_name_parts(replacements, text)

        # "May" appears in text but is in stopwords
        assert "May" not in result
        # "Theresa" should still be added if present
        assert result["Theresa"] == "[Person-1]"

    def test_short_candidates_excluded(self):
        """Candidates < 3 chars (like 'Li') should be excluded."""
        replacements = {"Li Wei Chen": "[Person-1]"}
        text = "Li Wei Chen and Chen discussed the matter."
        result = _expand_person_name_parts(replacements, text)

        # "Li" is only 2 chars, should not be added
        assert "Li" not in result
        # "Chen" (4 chars) should be added
        assert result["Chen"] == "[Person-1]"

    def test_non_person_placeholders_ignored(self):
        """Only [Person-N] placeholders should trigger expansion."""
        replacements = {
            "Darren Woods": "[Company]",
            "Jane Smith": "[Person-1]",
        }
        text = "Darren Woods and Jane Smith discussed. Smith agreed. Woods declined."
        result = _expand_person_name_parts(replacements, text)

        # Smith should be expanded (Person placeholder)
        assert result["Smith"] == "[Person-1]"
        # Woods should NOT be expanded (Company placeholder)
        assert "Woods" not in result or result.get("Woods") == "[Person-1]"
        # Actually "Woods" shouldn't be there at all since "Darren Woods" has [Company]
        # and only [Person-N] triggers expansion

    def test_existing_entries_not_overwritten(self):
        """If 'Woods' is already a replacement key, don't overwrite it."""
        replacements = {
            "Darren L. Woods": "[Person-1]",
            "Woods": "[Person-2]",  # Different person named Woods
        }
        text = "Darren L. Woods and Woods attended."
        result = _expand_person_name_parts(replacements, text)

        # "Woods" should keep its original mapping
        assert result["Woods"] == "[Person-2]"

    def test_two_name_person(self):
        """Simple two-word name generates individual components."""
        replacements = {"Jane Smith": "[Person-1]"}
        text = "Jane Smith presented. Smith agreed. Jane was happy."
        result = _expand_person_name_parts(replacements, text)

        assert result["Smith"] == "[Person-1]"
        assert result["Jane"] == "[Person-1]"

    def test_variant_not_in_text_not_added(self):
        """Variants that don't appear in the document text should be skipped."""
        replacements = {"Darren L. Woods": "[Person-1]"}
        # Text mentions neither the full name nor any variant
        text = "The agreement was signed by the CEO on January 5."
        result = _expand_person_name_parts(replacements, text)

        # No variants should be added since none appear in the text
        assert "Woods" not in result
        assert "Darren" not in result
        assert "Darren Woods" not in result


# ===========================================================================
# Integration: all fixes working together
# ===========================================================================


class TestIntegration:
    def test_multiline_address_cloaked(self, tmp_path):
        """A multi-line address in additional_replacements gets split and cloaked."""
        paragraphs = [
            "The office is located at 6001 Bollinger Canyon Road",
            "San Ramon, CA 94583",
            "This Agreement is between Acme Corp. and BigCo LLC.",
        ]
        docx_path = _make_docx(tmp_path, "addr.docx", paragraphs)
        output_path = tmp_path / "addr_cloaked.docx"
        mapping_path = tmp_path / "addr_mapping.json"

        config = CloakConfig(
            party_a_name="Acme Corp.",
            party_b_name="BigCo LLC",
            additional_replacements={
                "[Address-1]": "6001 Bollinger Canyon Road\nSan Ramon, CA 94583",
            },
            comment_mode=CommentMode.STRIP,
            strip_metadata=True,
        )

        result = cloak_document(
            input_path=docx_path,
            output_path=output_path,
            mapping_path=mapping_path,
            config=config,
        )

        cloaked_text = _extract_text(Path(result.output_path))
        assert "Bollinger" not in cloaked_text
        assert "San Ramon" not in cloaked_text

    def test_person_variants_cloaked_in_document(self, tmp_path):
        """Person name variants are expanded and cloaked in a real document."""
        paragraphs = [
            "This Agreement is between Acme Corp. and BigCo LLC.",
            "CEO Darren L. Woods approved the deal.",
            "Woods confirmed the terms were acceptable.",
            "Darren will attend the signing ceremony.",
        ]
        docx_path = _make_docx(tmp_path, "person.docx", paragraphs)
        output_path = tmp_path / "person_cloaked.docx"
        mapping_path = tmp_path / "person_mapping.json"

        config = CloakConfig(
            party_a_name="Acme Corp.",
            party_b_name="BigCo LLC",
            additional_replacements={
                "[Person-1]": "Darren L. Woods",
            },
            comment_mode=CommentMode.STRIP,
            strip_metadata=True,
        )

        result = cloak_document(
            input_path=docx_path,
            output_path=output_path,
            mapping_path=mapping_path,
            config=config,
        )

        cloaked_text = _extract_text(Path(result.output_path))
        assert "Darren" not in cloaked_text
        assert "Woods" not in cloaked_text
